"""
update_data.py — Dashboard OH · OM NRE (versi web)

Membaca tiga file Excel lalu menulis data.json untuk dashboard web:
  1. Jadwal OH NRE*.xlsx            -> jadwal  (sheet per tahun: 2023, 2024, ...)
  2. Jadwal OH OM 5*.xlsx           -> ditgas  (sheet "Run GPT AI", OM V)
  3. Renc&Real OH OM4*.xlsx         -> om4     (sheet "RencReal ...", OM IV)

data.json selalu DITIMPA (bukan ditambah), jadi ukurannya tetap.
Kalau isi data tidak berubah, file tidak ditulis ulang dan tidak di-push.
Kalau folder ini adalah repo git, perubahan otomatis di-commit dan di-push.

Cara pakai:  python update_data.py      (atau klik 2x update.bat)
"""
import datetime as dt
import glob
import json
import os
import re
import subprocess
import sys
import warnings

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl belum terpasang. Jalankan: pip install openpyxl")

warnings.filterwarnings("ignore")

# ---------------------------------------------------------------- pengaturan
HERE = os.path.dirname(os.path.abspath(__file__))
# Folder tempat file Excel. Default: folder induk dari folder repo ini.
# Bisa diganti, misalnya: EXCEL_DIR = r"C:\Users\Raheldan\OneDrive\Dokumen\claude database"
EXCEL_DIR = os.environ.get("OH_EXCEL_DIR") or os.path.dirname(HERE)

POLA_NRE = "Jadwal OH NRE*.xlsx"
POLA_OM5 = "Jadwal OH OM 5*.xlsx"
POLA_OM4 = "Renc&Real OH OM4*.xlsx"

# UP yang hanya dihitung di 2023 (sudah tidak masuk rekap NRE sejak 2024)
UP_KHUSUS_2023 = ("Kendari", "Pekanbaru")

OUT = os.path.join(HERE, "data.json")


# ---------------------------------------------------------------- utilitas
def cari_file(pola):
    files = [f for f in glob.glob(os.path.join(EXCEL_DIR, pola)) if not os.path.basename(f).startswith("~$")]
    if not files:
        return None
    return max(files, key=os.path.getmtime)


def iso(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) and v.year > 2000 else None


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def rapi(v):
    return " ".join(str(v).split()) if v is not None and str(v).strip() else None


# ---------------------------------------------------------------- 1. Jadwal OH NRE
def baca_nre(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    tahun_sheets = sorted(int(ws.title) for ws in wb if re.fullmatch(r"\d{4}", ws.title.strip()))
    out = []
    for y in tahun_sheets:
        ws = wb[str(y)]
        hdr = {}
        for c in ws[2]:
            if c.value is None:
                continue
            hdr.setdefault(str(c.value).strip().split("\n")[0], []).append(c.column - 1)

        def col(*names):
            for n in names:
                if n in hdr:
                    return hdr[n][0]
            return None

        cJP, cU, cMW = col("Jenis Pembangkit"), col("Unit Pembangkit"), col("DMN (MW)")
        cT, cSc, cD = col("Rencana/Realisasi", "Status"), col("Scope OH"), col("Durasi")
        cS, cE, cB, cP = col("Mulai"), col("Selesai"), col("Biaya"), col("Progress")
        cSt = [i for i in hdr.get("Status", []) if i != cT]
        cSt = cSt[0] if cSt else None
        cRP = {k: col(k) for k in ["R1", "R2", "R3", "P1", "P2", "P3"]}
        if cT is None or cS is None or cU is None:
            continue
        last, orph = {}, {}
        for r in ws.iter_rows(min_row=4, values_only=True):
            g = lambda i: r[i] if i is not None and i < len(r) else None
            t = g(cT)
            if t not in ("Rencana", "Realisasi") or not isinstance(g(cS), dt.datetime):
                continue
            up = (g(0) or "").strip()
            if y > 2023 and any(k in up for k in UP_KHUSUS_2023):
                continue
            st = g(cSt)
            st = str(st).strip() if st not in (None, " ", "") else None
            jp = g(cJP).strip() if isinstance(g(cJP), str) and g(cJP).strip() else None
            d = dict(up=up.replace("UPDK ", "UP "), jp=jp, u=str(g(cU)).strip(), mw=num(g(cMW)) if not isinstance(g(cMW), str) else None,
                     sc=str(g(cSc) or "").strip(), dur=num(g(cD)) if not isinstance(g(cD), str) else None,
                     s=iso(g(cS)), e=iso(g(cE)), biaya=g(cB) if isinstance(g(cB), (int, float)) else None,
                     st=st, p=g(cP) if isinstance(g(cP), (int, float)) else None,
                     rp={k.lower(): iso(g(i)) for k, i in cRP.items()})
            if not d["jp"]:
                for k in ["PLTMH", "PLTM", "PLTA", "PLTD", "PLTG"]:
                    if d["u"].startswith(k):
                        d["jp"] = k
                        break
            if t == "Rencana":
                d["year"] = y
                d["real"] = orph.pop(d["u"], None)
                out.append(d)
                last[d["u"]] = d
            else:
                pl = last.get(d["u"])
                if pl and pl["real"] is None:
                    pl["real"] = d
                else:
                    orph[d["u"]] = d

    def kat(x):
        R = x["real"]
        if R:
            s = (R["st"] or "").upper()
            if "GOING" in s:
                return "On going"
            if "EXTEND" in s:
                return "Extend"
            if "OPTIM" in s:
                return "Optimasi"
            return "On time"
        s = (x["st"] or "").lower()
        if "cancel" in s or "batal" in s:
            return "Batal"
        if "force" in s:
            return "Force majeure"
        if any(k in s for k in ["pending", "diundur", "geser", "akhir juli"]):
            return "Pending / diundur"
        return "Belum terlaksana"

    rows, cnt = [], {}
    for x in out:
        y = x["year"]
        cnt[y] = cnt.get(y, 0) + 1
        R = x["real"] or {}
        rows.append(dict(id=f"{y}-{cnt[y]:03d}", tahun=y, up=x["up"], jenis=x["jp"], unit=x["u"], dmn_mw=x["mw"], scope=x["sc"],
                         rencana_mulai=x["s"], rencana_selesai=x["e"], rencana_hari=x["dur"], rencana_biaya_ribu=x["biaya"],
                         keterangan=x["st"] if not x["real"] else None,
                         real_mulai=R.get("s"), real_selesai=R.get("e"), real_hari=R.get("dur"), real_biaya_ribu=R.get("biaya"),
                         status_realisasi=R.get("st"), progress=R.get("p"), kategori=kat(x), **x["rp"]))
    return rows


# ---------------------------------------------------------------- 2. OM V (Ditgas)
def baca_om5(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Run GPT AI"]
    upd = iso(ws["J2"].value)
    rows, n = [], 0
    for r in ws.iter_rows(min_row=5, values_only=True):
        r = list(r) + [None] * 20
        up, unit, st = r[2], r[3], r[8]
        if not unit or not st or not up:
            continue
        st = str(st).strip()
        n += 1
        s = st.lower()
        kat = {"real": "Selesai", "progress": "On going", "renc": "Rencana", "meluncur": "Meluncur 2027", "cancel": "Batal"}.get(s)
        if kat is None:
            kat = "Rencana revisi" if s.startswith("rev") else st
        rows.append(dict(id=f"dg-{n:03d}", up=rapi(up), unit=rapi(unit), sn=rapi(r[4]), dmn_mw=num(r[5]),
                         status_file=st, kategori=kat, scope=rapi(r[9]),
                         durasi_rencana=num(r[10]), durasi_real=num(r[11]), mulai=iso(r[12]), selesai_rencana=iso(r[13]),
                         selesai_real=iso(r[14]), progress=num(r[15]), selisih_hari=num(r[16]),
                         ketepatan=rapi(r[17]), info=rapi(r[18]), data_per=upd))
    return rows


# ---------------------------------------------------------------- 3. OM IV
def baca_om4(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = next((w for w in wb if w.title.lower().startswith("rencreal")), wb.worksheets[0])
    m = re.search(r"(\d{4})[.\-_](\d{2})[.\-_](\d{2})", os.path.basename(path))
    data_per = f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else dt.date.fromtimestamp(os.path.getmtime(path)).isoformat()
    R = [list(r) + [None] * 16 for r in ws.iter_rows(min_row=5, max_col=16, values_only=True)]
    out, n = [], 0
    for i, r in enumerate(R):
        if r[5] != "Renc" or not r[3]:
            continue
        n += 1
        nx = R[i + 1] if i + 1 < len(R) else [None] * 16
        is_real = nx[5] == "Real" or (nx[3] is None and nx[0] is None and (iso(nx[10]) or isinstance(nx[12], (int, float))))
        mulai, selesai, rot = iso(r[10]), iso(r[11]), False
        if not mulai and iso(r[14]):
            mulai, selesai, rot = iso(r[14]), iso(r[15]), True
        dmn = num(r[4])
        if dmn and dmn > 1000:  # tertulis dalam kW
            dmn = dmn / 1000
        dur = num(r[9]) or num(r[13])
        rr = dict(dur=None, s=None, e=None, p=None, kat=None, note=None)
        if is_real:
            rr = dict(dur=round(num(nx[9]), 1) if num(nx[9]) is not None else None, s=iso(nx[10]), e=iso(nx[11]),
                      p=num(nx[12]), kat=nx[7], note=rapi(nx[6]))
        p = rr["p"]
        if rr["s"] and p is not None and p >= 1:
            kat = "Selesai"
        elif rr["s"]:
            kat = "On going"
        else:
            kat = "Rencana"
        sel = tep = None
        if kat == "Selesai" and rr["e"] and selesai:
            sel = (dt.date.fromisoformat(rr["e"]) - dt.date.fromisoformat(selesai)).days
            tep = "Lead" if sel < 0 else ("On Time" if sel == 0 else "Extend")
        info = []
        if rot:
            info.append("Jadwal mengikuti ROT (origin schedule)")
        if rr["note"]:
            info.append(rr["note"].strip("()"))
        if rr["kat"]:
            info.append("Kategori " + str(rr["kat"]))
        out.append(dict(id=f"om4-{n:03d}", up="UP " + rapi(r[1]).title(), unit=rapi(r[3]), jenis=r[2], sn=None, dmn_mw=dmn,
                        status_file="Real" if kat != "Rencana" else "Renc", kategori=kat, scope=rapi(r[6]),
                        durasi_rencana=dur, durasi_real=rr["dur"], mulai=mulai, selesai_rencana=selesai,
                        selesai_real=rr["e"] if kat == "Selesai" else None, progress=p if rr["s"] else None,
                        selisih_hari=sel, ketepatan=tep, info="; ".join(info) or None, data_per=data_per))
    return out


# ---------------------------------------------------------------- utama
def git(*args):
    return subprocess.run(["git", *args], cwd=HERE, capture_output=True, text=True)


def main():
    print(f"Folder Excel: {EXCEL_DIR}")
    data, sumber = {}, {}
    for key, pola, fn in [("jadwal", POLA_NRE, baca_nre), ("ditgas", POLA_OM5, baca_om5), ("om4", POLA_OM4, baca_om4)]:
        f = cari_file(pola)
        if not f:
            print(f"  ! {pola} tidak ditemukan, data {key} lama dipakai")
            continue
        try:
            data[key] = fn(f)
        except Exception as e:  # noqa: BLE001
            print(f"  ! Gagal membaca {os.path.basename(f)}: {e}")
            continue
        sumber[key] = dict(file=os.path.basename(f), diubah=dt.datetime.fromtimestamp(os.path.getmtime(f)).strftime("%Y-%m-%d %H:%M"))
        print(f"  ok {os.path.basename(f)}: {len(data[key])} baris")

    lama = {}
    if os.path.exists(OUT):
        with open(OUT, encoding="utf-8") as fh:
            lama = json.load(fh)
    for key in ("jadwal", "ditgas", "om4"):
        if key not in data and key in lama:
            data[key] = lama[key]
            sumber[key] = lama.get("sumber", {}).get(key)

    if all(lama.get(k) == data.get(k) for k in ("jadwal", "ditgas", "om4")):
        print("Tidak ada perubahan data. data.json tidak diubah.")
        return

    baru = dict(dibuat=dt.datetime.now().strftime("%Y-%m-%d %H:%M"), sumber=sumber, **data)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(baru, fh, ensure_ascii=False, separators=(",", ":"))
    print(f"data.json ditulis ({os.path.getsize(OUT) // 1024} KB)")

    if os.path.isdir(os.path.join(HERE, ".git")):
        git("add", "data.json")
        c = git("commit", "-m", f"Update data {baru['dibuat']}")
        if c.returncode != 0:
            print(c.stdout or c.stderr)
            return
        p = git("push")
        print("Push ke GitHub berhasil." if p.returncode == 0 else f"Push gagal:\n{p.stderr}")
    else:
        print("Folder ini belum repo git, jadi tidak di-push.")


if __name__ == "__main__":
    main()
